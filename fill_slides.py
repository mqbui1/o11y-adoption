"""
Fill in placeholder content for slides 1-9 of the HOPe deck.
Preserves all existing formatting — only replaces text content.
"""
from pptx import Presentation
import copy

SRC = '/Users/mbui/Documents/o11y-adoption/reports/Observability Health Checks - MB.pptx'
DST = '/Users/mbui/Documents/o11y-adoption/reports/Observability Health Checks - Filled.pptx'

prs = Presentation(SRC)


def set_para_text(para, text):
    """Replace paragraph text while preserving the first run's formatting."""
    if not para.runs:
        para.add_run().text = text
        return
    para.runs[0].text = text
    for run in para.runs[1:]:
        run.text = ''


def set_tf_bullets(tf, bullets):
    """Set text frame paragraphs to a list of strings. Reuses/extends existing paragraphs."""
    paras = tf.paragraphs
    for i, text in enumerate(bullets):
        if i < len(paras):
            set_para_text(paras[i], text)
        else:
            new_para = copy.deepcopy(paras[-1]._p)
            tf._txBody.append(new_para)
            set_para_text(tf.paragraphs[i], text)
    for i in range(len(bullets), len(tf.paragraphs)):
        set_para_text(tf.paragraphs[i], '')


def shape_by_id(slide, shape_id):
    for s in slide.shapes:
        if s.shape_id == shape_id:
            return s
    return None


def set_cell(table, row, col, text):
    cell = table.rows[row].cells[col]
    for para in cell.text_frame.paragraphs:
        set_para_text(para, '')
    set_para_text(cell.text_frame.paragraphs[0], text)


# ── Slide 1: Cover ───────────────────────────────────────────────────────────
# Title is a template instruction — replace with the real deck title
slide1 = prs.slides[0]
title = shape_by_id(slide1, 2)
if title:
    set_para_text(title.text_frame.paragraphs[0], 'HOPe — Helping O11y renewal Program')


# ── Slide 2: Program Overview ────────────────────────────────────────────────
slide2 = prs.slides[1]
table_shape = shape_by_id(slide2, 6259)
if table_shape and hasattr(table_shape, 'table'):
    tbl = table_shape.table

    # Overview & Objectives
    set_cell(tbl, 0, 1,
        'Provide SAs with standardised health check reports (o11y-adoption + o11y-usage-governance) '
        'to proactively identify at-risk O11y customers 12+ months before renewal. '
        'Goal: get ahead of churn signals before they become renewal losses.'
    )

    # Business Outcome
    set_cell(tbl, 1, 1,
        'Reduce O11y churn by surfacing leading indicators early — inactive users, stale assets, '
        'OTel coverage gaps, and Custom Metrics overages. '
        'Enable SAs to have data-driven renewal conversations and improve renewal rates for $500K+ customers.'
    )

    # Success Metrics
    set_cell(tbl, 2, 1,
        '• Adoption health score improvement across pilot customers (target: avg score >65/100)\n'
        '• Reduction in Custom Metrics overages flagged post-review\n'
        '• % of at-risk $500K+ O11y renewals with health check completed 12+ months out\n'
        '• Renewal rate improvement vs. baseline for customers receiving health checks'
    )


# ── Slide 3: Owners & Stakeholders ──────────────────────────────────────────
slide3 = prs.slides[2]
replacements_s3 = {
    20: 'Marc Bui',
    21: 'Joe Starofsky',
    22: 'TBD',
    23: 'Program Lead — SA/CSE',
    24: 'Executive Sponsor — O11y GTM',
    25: 'Project Manager',
    33: 'SA Enablement',
    34: 'O11y Product Management',
    35: 'Field Marketing',
    36: 'SA Enablement Lead',
    37: 'O11y PM',
    38: 'Marketing Partner',
    48: 'Regional SA Leads',
    49: 'Renewal Managers',
    50: 'Customer Success',
    51: 'SA Leadership',
    52: 'CS Leadership',
    53: 'Renewal Management',
    62: 'Finance',
    63: 'Legal/Compliance',
    64: 'Product Analytics',
    65: 'FP&A',
    66: 'Trust & Safety',
    67: 'Data Engineering',
}
for shape_id, text in replacements_s3.items():
    shape = shape_by_id(slide3, shape_id)
    if shape and shape.has_text_frame:
        set_para_text(shape.text_frame.paragraphs[0], text)


# ── Slide 4: Current State ───────────────────────────────────────────────────
slide4 = prs.slides[3]
# Fix "Subtitle" heading and expand implication bullet
shape7 = shape_by_id(slide4, 7)
if shape7:
    set_tf_bullets(shape7.text_frame, [
        'Involvement before its too late',
        'Loss of millions in revenue from churning & struggling customers.',
        'Soft costs including demos, enablement & POCs to replace lost revenue.',
        'Missed signals that o11y-adoption and o11y-usage-governance reports would have surfaced early.',
    ])


# ── Slide 5: Proposed Future State ──────────────────────────────────────────
slide5 = prs.slides[4]

shape5 = shape_by_id(slide5, 5)
if shape5:
    set_tf_bullets(shape5.text_frame, [
        'Frequent Health Reviews!',
        'Focus on customers $500K and greater.',
        'Early engagements — 12+ months before O11y renewal.',
        'Run o11y-adoption + o11y-usage-governance reports as a standard SA deliverable.',
    ])

shape6 = shape_by_id(slide5, 6)
if shape6:
    set_tf_bullets(shape6.text_frame, [
        'Health Reviews',
        'Look at 12+ months out for O11y Renewals.',
        'Send frequent health review reports with prioritised actions to optimise customer environment.',
        'Use o11y-adoption report to surface inactive users, stale assets, and OTel coverage gaps.',
        'Use o11y-usage-governance report to identify cardinality explosions and Custom Metrics overages.',
        'EARLIER engagements — before renewal conversations begin.',
    ])

shape7_s5 = shape_by_id(slide5, 7)
if shape7_s5:
    set_tf_bullets(shape7_s5.text_frame, [
        'Involvement before its too late',
        'Identify positive & negative adoption trends of the customer.',
        'Call out potential licensing issues such as Custom Metrics overages before they become billing surprises.',
        'Improve renewal rates by demonstrating platform value early.',
        'Reduce churn risk with data-driven health scores and recommended actions.',
    ])


# ── Slide 6: Program Status ──────────────────────────────────────────────────
slide6 = prs.slides[5]

accomplishments = shape_by_id(slide6, 10)
if accomplishments:
    set_tf_bullets(accomplishments.text_frame, [
        'o11y-adoption report built and validated against real org data (71/100, B grade).',
        'o11y-usage-governance cardinality report built and validated.',
        'Both tools work with "usage" API role — minimal permissions required.',
        'HTML reports are self-contained and shareable with no infrastructure needed.',
    ])

in_progress = shape_by_id(slide6, 6)
if in_progress:
    set_tf_bullets(in_progress.text_frame, [
        'Refining report content — needs team review to confirm data is meaningful and actionable for customers.',
        'Automating report generation and delivery at a regular cadence as a customer-facing slide deck.',
    ])

risks = shape_by_id(slide6, 11)
if risks:
    set_tf_bullets(risks.text_frame, [
        'Customer org access — customers may be prohibited from sharing an API token into their org.',
        'Data privacy — report surfaces org-level user and asset data; legal review needed before sharing externally.',
        'SA bandwidth — report runs are manual today; no automated scheduling yet.',
        'Audit API limitations — no read/view tracking; engagement inferred from write activity only.',
    ])


# ── Slide 7: Leadership Update ───────────────────────────────────────────────
# Already fully populated with real content — no changes needed


# ── Slide 8: Project Roadmap ─────────────────────────────────────────────────
slide8 = prs.slides[7]

# Dates
date_shapes = {3: 'Q2 2026', 4: 'Q3 2026', 9: 'Q4 2026+'}
for shape_id, text in date_shapes.items():
    shape = shape_by_id(slide8, shape_id)
    if shape and shape.has_text_frame:
        set_para_text(shape.text_frame.paragraphs[0], text)

# Phase bullet content (shapes 22, 23, 24)
phase_bullets = {
    22: [
        'Run o11y-adoption + o11y-usage-governance reports on 3-5 at-risk customers ($500K+) manually via SA.',
        'Validate report accuracy, confirm token/API access requirements, and identify content gaps.',
    ],
    23: [
        'Standardise report cadence (monthly/quarterly) for all at-risk O11y renewals 12+ months out.',
        'Train SA team on interpreting and presenting reports to customers.',
        'Define report delivery format and customer-facing narrative.',
    ],
    24: [
        'Integrate reports into renewal playbook as a required pre-renewal touchpoint.',
        'Automate report delivery and trending via scheduled runs.',
        'Explore hosting in Dogfood or equivalent internal tool for broader SA access.',
    ],
}
for shape_id, bullets in phase_bullets.items():
    shape = shape_by_id(slide8, shape_id)
    if shape and shape.has_text_frame:
        set_tf_bullets(shape.text_frame, bullets)


# ── Slide 9: Status Dashboard ────────────────────────────────────────────────
slide9 = prs.slides[8]

# Upcoming Tasks (shape 13)
tasks = shape_by_id(slide9, 13)
if tasks:
    set_tf_bullets(tasks.text_frame, [
        'Identify 3-5 pilot customers ($500K+ at-risk O11y renewals)',
        'Get customer org API token access for pilot accounts',
        'Run adoption + cardinality reports for pilot customers and review with SA team',
    ])

# Completed (shape 19)
completed = shape_by_id(slide9, 19)
if completed:
    set_tf_bullets(completed.text_frame, [
        'o11y-adoption report built and validated',
        'o11y-usage-governance report built and validated',
        'Both tools tested with usage API role',
    ])

# Key Milestones table (shape 10)
milestones = shape_by_id(slide9, 10)
if milestones and hasattr(milestones, 'table'):
    set_cell(milestones.table, 0, 0, 'o11y-adoption + cardinality tools built')
    set_cell(milestones.table, 0, 1, 'Complete')
    set_cell(milestones.table, 1, 0, 'Pilot customer list identified')
    set_cell(milestones.table, 1, 1, 'In Progress')
    set_cell(milestones.table, 2, 0, 'SA team trained on report delivery')
    set_cell(milestones.table, 2, 1, 'Planned')

# Risks & Blockers table (shape 16)
risks_table = shape_by_id(slide9, 16)
if risks_table and hasattr(risks_table, 'table'):
    set_cell(risks_table.table, 1, 0, 'Customer org API token access')
    set_cell(risks_table.table, 1, 1, 'Work with SA/CSE to obtain usage-role token per customer')
    set_cell(risks_table.table, 1, 2, 'In Progress')
    set_cell(risks_table.table, 2, 0, 'SA bandwidth for manual report runs')
    set_cell(risks_table.table, 2, 1, 'Automate scheduling; scope pilot to 3-5 customers initially')
    set_cell(risks_table.table, 2, 2, 'Planning')
    set_cell(risks_table.table, 3, 0, 'Customer data privacy')
    set_cell(risks_table.table, 3, 1, 'Review with Legal before sharing reports externally')
    set_cell(risks_table.table, 3, 2, 'Planning')


prs.save(DST)
print(f'Saved: {DST}')
